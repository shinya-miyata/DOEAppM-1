import streamlit as st
import pandas as pd
import numpy as np
from sklearn.preprocessing import StandardScaler, PolynomialFeatures
from sklearn.linear_model import LinearRegression
from sklearn.pipeline import Pipeline
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import r2_score, mean_squared_error
import io

st.set_page_config(page_title="軽量DOEアプリ", layout="wide")
st.title("🔬 軽量版 実験計画支援アプリ")

# =========================================================
# ① 実験候補生成
# =========================================================
st.header("① 実験候補の生成")

uploaded_csv = st.file_uploader("生成条件CSV（1行目:項目名、2行目:lower、3行目:upper）", type="csv")
sample_size = st.number_input("生成するサンプル数", min_value=10, max_value=5000, value=100)

if uploaded_csv:
    df_bounds = pd.read_csv(uploaded_csv, index_col=0)
    df_bounds = df_bounds.dropna(axis=1, how="all")
    st.write("読み込んだ生成条件:", df_bounds)

    samples = []
    for _ in range(sample_size):
        row = []
        for col in df_bounds.columns:
            lower = float(df_bounds.loc["lower", col])
            upper = float(df_bounds.loc["upper", col])
            row.append(np.random.uniform(lower, upper))
        samples.append(row)

    df_samples = pd.DataFrame(samples, columns=df_bounds.columns)
    st.dataframe(df_samples)

    # 保存
    st.session_state["candidates_raw"] = df_bounds.copy()
    st.session_state["generated_candidates"] = df_samples.copy()

# =========================================================
# ② D最適基準で候補選択
# =========================================================
st.header("② 実験候補選択（D最適基準）")

uploaded_samples = st.file_uploader("候補データCSV（説明変数のみ）", type="csv")
select_size = st.number_input("選択するサンプル数", min_value=2, max_value=500, value=10)

if uploaded_samples:
    df_all = pd.read_csv(uploaded_samples)
    st.write("候補データ:", df_all)

    best_score = -np.inf
    best_subset = None

    for _ in range(200):
        subset = df_all.sample(n=select_size)
        X = StandardScaler().fit_transform(subset.values)
        score = np.linalg.det(X.T @ X)
        if score > best_score:
            best_score = score
            best_subset = subset

    df_unselected = df_all.drop(best_subset.index)

    st.subheader("選択された候補")
    st.dataframe(best_subset)

    st.subheader("未選択候補")
    st.dataframe(df_unselected)

    # 保存
    st.session_state["selected_candidates"] = best_subset.copy()
    st.session_state["unselected_candidates"] = df_unselected.copy()

# =========================================================
# ③ 次の候補選定（複数モデル対応）
# =========================================================
st.header("③ 次の実験候補選定（複数モデル対応）")

uploaded_train = st.file_uploader("実験済みデータ（説明変数＋目的変数）CSV", type="csv")
uploaded_candidates = st.file_uploader("未実験候補（説明変数のみ）CSV", type="csv")

if uploaded_train and uploaded_candidates:
    df_train = pd.read_csv(uploaded_train)
    df_candidates = pd.read_csv(uploaded_candidates)

    target_col = st.selectbox("目的変数を選択", df_train.columns)
    goal_value = st.number_input("目標値", value=float(df_train[target_col].median()))

    model_name = st.selectbox(
        "使用する回帰モデルを選択",
        ["LinearRegression", "PolynomialRegression", "RandomForest"]
    )

    def build_model(name):
        if name == "LinearRegression":
            return LinearRegression()
        elif name == "PolynomialRegression":
            return Pipeline([
                ("poly", PolynomialFeatures(degree=2)),
                ("linear", LinearRegression())
            ])
        elif name == "RandomForest":
            return RandomForestRegressor(n_estimators=200, random_state=0)

    model = build_model(model_name)

    feature_cols = [col for col in df_train.columns if col != target_col]

    # ★ 修正済み：feature_candidates → feature_cols
    missing_cols = [col for col in feature_cols if col not in df_candidates.columns]
    if missing_cols:
        st.error(f"候補データに不足している列があります: {missing_cols}")
    else:
        X = df_train[feature_cols].values
        y = df_train[target_col].values
        X_pred = df_candidates[feature_cols].values

        scaler = StandardScaler()
        X_scaled = scaler.fit_transform(X)
        X_pred_scaled = scaler.transform(X_pred)

        model.fit(X_scaled, y)

        y_pred_train = model.predict(X_scaled)
        r2 = r2_score(y, y_pred_train)
        rmse = np.sqrt(mean_squared_error(y, y_pred_train))

        st.subheader("📊 モデルの正確性")
        st.write(f"R²（決定係数）: **{r2:.3f}**")
        st.write(f"RMSE（誤差）: **{rmse:.3f}**")

        y_pred = model.predict(X_pred_scaled)

        df_result = df_candidates.copy()
        df_result["Predicted"] = y_pred
        df_result["Distance"] = np.abs(y_pred - goal_value)

        df_sorted = df_result.sort_values("Distance")
        st.subheader("📌 次の候補（目標値に近い順）")
        st.dataframe(df_sorted)

        # 保存
        st.session_state["train_data"] = df_train.copy()
        st.session_state["target_columns"] = [target_col]
        st.session_state["target_goals"] = {target_col: goal_value}
        st.session_state["model_name"] = model_name
        st.session_state["model_scores"] = {
            target_col: {
                "r2_train": r2,
                "rmse_train": rmse
            }
        }
        st.session_state["next_candidates"] = df_sorted.copy()

# =========================================================
# ④ Excelレポート生成（横並び・openpyxl安定版）
# =========================================================
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

st.header("📘 最終レポート生成（Excel）")

required_keys = [
    "candidates_raw",
    "generated_candidates",
    "selected_candidates",
    "unselected_candidates",
    "train_data",
    "target_columns",
    "target_goals",
    "model_name",
    "model_scores",
    "next_candidates"
]

missing = [k for k in required_keys if k not in st.session_state]

if missing:
    st.warning(f"以下のデータが不足しています: {missing}")
else:
    if st.button("📥 Excelレポートを生成する"):

        # ★ openpyxl で新規 Workbook を作成（安全）
        wb = Workbook()
        ws = wb.active
        ws.title = "DOEレポート"

        # セクションをまとめる
        sections = []

        def add_section(title, df=None, description=None):
            sections.append({"title": title, "description": description, "df": df})

        add_section("① 読み込んだ候補データ（lower/upper）",
                    st.session_state["candidates_raw"],
                    "候補生成に使用した下限値・上限値")

        add_section("② 生成された候補データ",
                    st.session_state["generated_candidates"],
                    "①の範囲からランダム生成した候補")

        add_section("③ D最適基準で選択された候補",
                    st.session_state["selected_candidates"],
                    "D最適基準により選ばれた候補")

        add_section("④ 未選択候補",
                    st.session_state["unselected_candidates"],
                    "選択されなかった候補")

        add_section("⑤ 読み込んだ実験済みデータ",
                    st.session_state["train_data"],
                    "説明変数と目的変数を含む実験済みデータ")

        df_goals = pd.DataFrame({
            "目的変数": st.session_state["target_columns"],
            "目標値": [st.session_state["target_goals"][col] for col in st.session_state["target_columns"]]
        })
        add_section("⑥ 目的変数と目標値",
                    df_goals,
                    "予測モデルが目指す目標値")

        df_model = pd.DataFrame({"使用モデル": [st.session_state["model_name"]]})
        add_section("⑦ 使用した回帰モデル",
                    df_model,
                    "予測に使用したモデル名")

        df_scores = pd.DataFrame(st.session_state["model_scores"]).T
        add_section("⑧ モデルの正確性（R², RMSE）",
                    df_scores,
                    "学習データに対するモデル性能")

        add_section("⑨ 次の実験候補（目標値に近い順）",
                    st.session_state["next_candidates"],
                    "予測値と目標値の距離が近い順に並べた候補")

        # ★ 横方向に貼り付ける
        col = 1

        for sec in sections:
            ws.cell(row=1, column=col, value=sec["title"])
            ws.cell(row=2, column=col, value=sec["description"])

            # DataFrame → Excel rows
            for r_idx, row in enumerate(dataframe_to_rows(sec["df"], index=False, header=True), start=3):
                for c_idx, value in enumerate(row, start=col):
                    ws.cell(row=r_idx, column=c_idx, value=value)

            col += len(sec["df"].columns) + 3  # 次のデータを右へ

        # ★ バイトデータとして保存
        output = io.BytesIO()
        wb.save(output)

        st.download_button(
            label="📄 Excelレポート（横並び）をダウンロード",
            data=output.getvalue(),
            file_name="DOE_Final_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
